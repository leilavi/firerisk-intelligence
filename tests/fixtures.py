"""TEST FIXTURES ONLY - small workbooks with the same column layout as the real sources, used to test
pipeline behaviour (duplicates, invalid values, missing tabs...). They are labelled as test data, are never
shipped as product data and never shown in the dashboard."""

from __future__ import annotations

import json
from datetime import datetime, time, timedelta
from pathlib import Path

import openpyxl

METER_COLS = ["Log_ID", "Equipment_ID", "Hình_ảnh_Thiết_Bị", "Tên_Hệ_Thống", "Branch", "Ngay_Kiem_Tra",
              "Thời_Điểm_Kiểm_Tra", "Ca_Truc", "Trạng_Thái_CV", "Hướng_Dẫn_Chung", "Hướng_Dẫn_Chi_Tiết", "Hình_ảnh_1",
              "Hình_ảnh_2", "Hình_ảnh_3", "Hình_ảnh_4", "Chỉ_Số_Thiết_Bị", "Tiêu_Thụ_Ngày", "Tiêu_Thụ_Lũy_Tiến",
              "Chỉ_Số_Bất_Thường", "Báo_Cáo_Bất_Thường", "Hình_ảnh_Chạy_Thử", "Thời_Điểm_Chạy_Thử",
              "Thời_Gian_Chạy_Thử", "Đánh_Giá_Cảm_Quan", "Tình_Trạng_Thiết_Bị", "Lý_Do_Cần_Theo_Dõi", "Hướng_Xử_Lý"]
ASSET_COLS = ["Log_ID", "Equipment_ID", "Location_ID", "Service Type", "Service_Details", "Maintenance_Details",
              "Current_Status", "Latest_Activity_Date", "Notify Entire Company", "Requester", "Approving_Mgr",
              "Estimated_Cost", "Estimated Duration (days)", "Expected_Completion_Date", "Approval_Status",
              "Mgr_Comment", "Maintenance_Date", "Technician_Name", "Action_Type", "Actual_Cost"]
INC_COLS = ["Mã sự cố/ nâng cấp", "Loại báo cáo", "Thời gian báo cáo", "Báo cáo nhanh", "Mô tả chi tiết",
            "Mức độ khẩn cấp", "Báo toàn hệ thống?", "Hướng xử lý, khắc phục", "Đơn vị đầu tư", "Trạng thái",
            "Chi Nhánh", "Khu/ Block", "Tầng", "Khu vực", "Phòng/ Unit", "Video_link_Url", "Ghi chú nhanh về Video",
            "Hình ảnh 1", "Ghi chú H1", "Người báo cáo", "Tg bắt đầu TC", "Tg Hoàn thành"]
GEN_COLS = ["ID", "Ngày_Vận_Hành", "Thiết_Bị_Vận_Hành", "Branch", "Chỉ_số_trước_VH", "Tg_bắt_đầu_VH",
            "Tg_kết_thúc_VH", "Thời_lượng_VH", "Chỉ-số_sau_VH", "Tiêu_hao_nhiên_liệu", "Lý_do_Vận_Hành",
            "Nhập_Nhiên_liệu", "Nhiên_liệu_tồn", "%_Nhiên_liệu_tồn", "Nhân_sự_VH", "Giá_NL_TB"]

NOW = datetime(2026, 10, 27, 14, 0)


def _row(cols, **kv):
    return [kv.get(c) for c in cols]


def build_ops_workbook(path: Path, *, drop_tab: str | None = None) -> Path:
    from firerisk.config import BRANCH_PREFIX as B
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    tabs = {}
    ws = wb.create_sheet("chỉ số điện nước"); tabs["m"] = ws
    ws.append(METER_COLS)
    for i in range(20):
        d = NOW - timedelta(days=20 - i)
        ws.append(_row(METER_COLS, Log_ID=f"m{i}", Equipment_ID="P-BG.TTC226", Branch=B,
                       Thời_Điểm_Kiểm_Tra=d.replace(hour=6), **{"Chỉ_Số_Thiết_Bị": 1000 + i * 9}))
    ws.append(_row(METER_COLS, Log_ID="mbad", Equipment_ID="P-BG.TTC172", Branch=B,
                   Thời_Điểm_Kiểm_Tra=NOW.replace(hour=6), **{"Chỉ_Số_Thiết_Bị": "đồng hồ hỏng"}))
    ws.append(_row(METER_COLS, Log_ID="C-tuBC", Equipment_ID="P-BG.TTC009", Branch=B,
                   Thời_Điểm_Kiểm_Tra=NOW.replace(hour=6), **{"Tình_Trạng_Thiết_Bị": "Đạt"}))
    ws = wb.create_sheet("BC hư hỏng TS"); ws.append(ASSET_COLS)
    ws.append(_row(ASSET_COLS, Log_ID="A1", Equipment_ID="P-BG.TTC167", Location_ID=f"{B}-A-B1-Chung-",
                   Service_Details="Máy bơm thổi khí kẹt cứng", Current_Status="2. Pending Repair",
                   Latest_Activity_Date=NOW - timedelta(days=6), Expected_Completion_Date=NOW - timedelta(days=3)))
    ws.append(_row(ASSET_COLS, Log_ID="A1", Equipment_ID="P-BG.TTC167", Location_ID=f"{B}-A-B1-Chung-",
                   Service_Details="Duplicate ID row", Current_Status="Trạng thái lạ",
                   Latest_Activity_Date=NOW - timedelta(days=5)))
    ws.append(_row(ASSET_COLS, Log_ID="Q1", Equipment_ID="E-ML1", Location_ID="SITE-2-B-F3",
                   Service_Details="other branch", Current_Status="4. Operational"))
    ws = wb.create_sheet("NK thi công - sự cố"); ws.append(INC_COLS)
    ws.append(_row(INC_COLS, **{"Mã sự cố/ nâng cấp": f"{B}----T10_1", "Loại báo cáo": "Báo cáo sự cố",
                                "Thời gian báo cáo": NOW - timedelta(days=19), "Báo cáo nhanh": "Phòng MBA,MSB thấm nước",
                                "Trạng thái": "3. Đang tìm PA", "Chi Nhánh": B}))
    ws.append(_row(INC_COLS, **{"Mã sự cố/ nâng cấp": f"{B}----T08_2", "Loại báo cáo": "Nhật ký thi công",
                                "Thời gian báo cáo": NOW - timedelta(days=80), "Báo cáo nhanh": "Sơn tường",
                                "Trạng thái": "Đang thi công", "Chi Nhánh": B,
                                "Tg Hoàn thành": NOW - timedelta(days=70)}))
    ws = wb.create_sheet("Dầu M Phát điện"); ws.append(GEN_COLS)
    ws.append(_row(GEN_COLS, ID="G1", Ngày_Vận_Hành=NOW - timedelta(days=4), Thiết_Bị_Vận_Hành="P-BG.TTC017",
                   Branch=B, Tg_bắt_đầu_VH=time(8, 0), Tg_kết_thúc_VH=time(8, 17), Thời_lượng_VH=time(0, 0),
                   Lý_do_Vận_Hành="VH định kỳ", Nhiên_liệu_tồn=1760))
    if drop_tab:
        wb.remove(wb[drop_tab])
    wb.save(path)
    return path


def build_snapshot(dirpath: Path, **kw) -> Path:
    dirpath.mkdir(parents=True, exist_ok=True)
    build_ops_workbook(dirpath / "ops.xlsx", **kw)
    (dirpath / "manifest.json").write_text(json.dumps({"created_by": "TEST FIXTURE", "files": {
        "ops_data": {"title": "TEST fixture", "local_file": "ops.xlsx"},
        "vehicles": {"title": "TEST fixture vehicles", "local_file": "missing_vehicles.xlsx"}}}), encoding="utf-8")
    return dirpath
