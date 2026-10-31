"""Draft 'panel -> life-safety equipment' from OCR results. Synthetic OCR json (TEST strings), no real drawing."""

import importlib.util
import json
import socket
from pathlib import Path

from openpyxl import load_workbook

from firerisk.power_links import load


def _mod():
    spec = importlib.util.spec_from_file_location("draft", Path(__file__).resolve().parents[1] / "scripts" / "draft_power_links.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def t(text, x, y, score=0.98):
    return {"text": text, "score": score, "box": [x, y, x + 60, y + 20], "rotated": False}


def write_ocr(folder: Path):
    ocr = folder / "ocr"
    ocr.mkdir(parents=True)
    lighting = {"file": "TECH-DIAG_TEST_Exit_Lighting.pdf", "pages": [
        {"page": 1, "size": [4000, 3000], "texts": [
            t("MAT BANG TANG HAM", 3000, 2800), t("TANG HAM", 3000, 2700),
            t("DB-TESTB1", 1000, 1000), t("EXIT", 1050, 1060), t("EM", 1100, 1010), t("L1-03", 1080, 1040),
            t("EXIT", 2600, 2000),                                   # far from the panel label: weaker guess
            t("DEN THOAT HIEM EXIT LOAI 1 MAT", 3200, 200),          # legend: not counted
            t("L1-05", 300, 300), t("L1-06", 320, 330)]},            # circuit codes must not vote for floor 1
        {"page": 2, "size": [4000, 3000], "texts": [
            t("TANG 1", 3000, 2800), t("TANG 2", 3000, 2700),       # riser / legend page: floor not clear
            t("EXIT", 500, 500, score=0.62)]},
    ]}
    cable = {"file": "TECH-DIAG_TEST_Cable_MSB.pdf", "pages": [
        {"page": 1, "size": [4000, 3000], "texts": [
            t("TANG HAM", 3000, 2800), t("MSB", 500, 500), t("BOM CHUA CHAY", 540, 560), t("4x25mm XLPE", 900, 900)]}]}
    (ocr / "a.json").write_text(json.dumps(lighting), encoding="utf-8")
    (ocr / "b.json").write_text(json.dumps(cable), encoding="utf-8")


def test_draft_rows_floors_panels_legend_and_numbers_only_output(tmp_path, capsys, monkeypatch):
    m = _mod()
    write_ocr(tmp_path)
    monkeypatch.setattr(m, "load_assets", lambda: {"P-TESTDB01": "TEST lighting panel DB-TESTB1 basement",
                                                   "P-TESTMSB": "TEST main switchboard MSB"})
    monkeypatch.setattr(socket, "socket", lambda *a, **k: (_ for _ in ()).throw(AssertionError("network used")))
    assert m.main(["x", str(tmp_path)]) == 0
    out = capsys.readouterr().out
    assert "TESTB1" not in out and "XLPE" not in out and "L1-03" not in out          # numbers only
    wb = load_workbook(tmp_path / "power_links_draft.xlsx")
    head = [c.value for c in wb["Links"][1]]
    rows = [dict(zip(head, [c.value for c in r])) for r in wb["Links"].iter_rows(min_row=2)]
    by = {(r["floor"] or "", r["panel_label"] or "", r["feeds"]): r for r in rows}
    near_exit = by[("B1", "DB-TESTB1", "exit_sign")]
    assert near_exit["count"] in (None, "") and near_exit["auto_count_hint"] == 2   # never pre-filled
    assert "panel label next to the symbol (1)" in near_exit["basis"] and "weaker" in near_exit["basis"]
    assert near_exit["asset_id"] == "P-TESTDB01" and "L1-03" in near_exit["circuits"]
    assert by[("B1", "DB-TESTB1", "emergency_light")]["auto_count_hint"] == 1
    assert by[("B1", "MSB", "fire_pump")]["asset_id"] == "P-TESTMSB"
    unclear = by[("", "", "exit_sign")]
    assert "floor not clear" in unclear["basis"] and unclear["ocr_confidence"] == 0.62
    assert all(not r["confirmed"] for r in rows)                       # nothing confirmed automatically
    assert sum(r["auto_count_hint"] for r in rows if r["feeds"] == "exit_sign") == 3   # legend phrase not counted
    panels = {r[0].value: r for r in wb["Panels"].iter_rows(min_row=2)}
    assert set(panels) == {"DB-TESTB1", "MSB"}


def test_the_draft_is_unused_until_confirmed_and_round_trips_after(tmp_path, monkeypatch):
    m = _mod()
    write_ocr(tmp_path)
    monkeypatch.setattr(m, "load_assets", lambda: {"P-TESTDB01": "TEST DB-TESTB1"})
    m.main(["x", str(tmp_path)])
    draft = tmp_path / "power_links_draft.xlsx"
    assert load(draft).links == []                                     # draft as written: nothing used
    wb = load_workbook(draft)
    ws = wb["Links"]
    head = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2):
        d = dict(zip(head, r))
        if d["floor"].value == "B1" and d["panel_label"].value == "DB-TESTB1":
            d["confirmed"].value = "Y"
    wb.save(tmp_path / "power_links.xlsx")
    pl = load(tmp_path / "power_links.xlsx", known_assets={"P-TESTDB01"})
    assert pl.problems == [] and len(pl.links) == 2
    deps = {d["equipment"]: d["count_known"] for d in pl.dependents("P-TESTDB01")}
    assert deps == {"exit_sign": False, "emergency_light": False}       # count left empty = "number not recorded"


def test_page_floor_is_left_empty_when_two_floors_are_named_equally():
    m = _mod()
    assert m.page_floor([t("TANG 1", 0, 0), t("TANG 2", 0, 0)]) is None
    assert m.page_floor([t("TANG HAM", 0, 0), t("TANG HAM", 0, 0), t("TANG 1", 0, 0)]) == "B1"
    assert m.page_floor([t("L1-05", 0, 0), t("L1", 0, 0)]) is None               # circuit codes do not vote
    assert m.panel_label("tủ điện DB-B1") == "DB-B1" and m.panel_label("EXIT") is None
    assert m.device_type("EXIT") == "exit_sign" and m.device_type("DEN THOAT HIEM EXIT LOAI 1 MAT") is None
