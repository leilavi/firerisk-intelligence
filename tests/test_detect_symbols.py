"""Legend-based symbol search on a generated TEST drawing (no real drawing)."""

import importlib.util
import json
from pathlib import Path

import pytest

cv2 = pytest.importorskip("cv2")
import numpy as np  # noqa: E402


def _mod(name):
    spec = importlib.util.spec_from_file_location(name, Path(__file__).resolve().parents[1] / "scripts" / f"{name}.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


def em_symbol(img, x, y, rot=False):
    """Emergency light: two circles on a box (like the legend '00' over a rectangle)."""
    s = np.full((40, 70), 255, np.uint8)
    cv2.rectangle(s, (2, 20), (67, 37), 0, 2)
    cv2.circle(s, (20, 11), 9, 0, 2)
    cv2.circle(s, (50, 11), 9, 0, 2)
    if rot:
        s = np.rot90(s)
    h, w = s.shape
    img[y:y + h, x:x + w] = np.minimum(img[y:y + h, x:x + w], s)


def exit_symbol(img, x, y):
    s = np.full((34, 80), 255, np.uint8)
    cv2.rectangle(s, (1, 1), (78, 32), 0, 2)
    cv2.putText(s, "EXIT", (8, 25), cv2.FONT_HERSHEY_SIMPLEX, 0.7, 0, 2)
    img[y:y + 34, x:x + 80] = np.minimum(img[y:y + 34, x:x + 80], s)


def build(folder: Path):
    ocr = folder / "ocr"
    ocr.mkdir(parents=True)
    img = np.full((1800, 2600), 255, np.uint8)
    # legend: symbol, then its name to the right (the OCR box of the name is given below)
    em_symbol(img, 2000, 1500)
    exit_symbol(img, 1995, 1580)
    cv2.rectangle(img, (1960, 1460), (2580, 1660), 0, 2)                  # legend frame
    # the plan: 5 emergency lights (2 turned), 3 exit signs, distractors
    for x, y, r in [(200, 200, False), (600, 300, False), (900, 900, True), (1300, 400, False), (300, 1200, True)]:
        em_symbol(img, x, y, r)
    for x, y in [(700, 700), (1500, 1000), (400, 600)]:
        exit_symbol(img, x, y)
    cv2.circle(img, (1700, 300), 9, 0, 2)                                  # a lone circle is not a symbol
    for x in range(0, 2600, 260):
        cv2.line(img, (x, 0), (x, 1400), 0, 1)                             # grid lines across the plan
    cv2.imwrite(str(ocr / "TEST_p1.png"), img)
    texts = [{"text": "ĐÈN CHIẾU SÁNG SỰ CỐ", "score": 0.97, "box": [2120, 1505, 2450, 1530], "rotated": False},
             {"text": "ĐÈN THOÁT HIỂM", "score": 0.97, "box": [2120, 1585, 2380, 1610], "rotated": False},
             {"text": "TANG HAM", "score": 0.97, "box": [100, 1700, 300, 1730], "rotated": False}]
    (ocr / "TECH-DIAG_TEST_Exit_Lighting.json").write_text(json.dumps(
        {"file": "TECH-DIAG_TEST_Exit_Lighting.pdf", "pages": [{"page": 1, "size": [2600, 1800], "image": "TEST_p1.png",
                                                                "texts": texts}]}), encoding="utf-8")
    return ocr


def test_symbols_are_cut_from_the_legend_and_counted_on_the_plan_only(tmp_path):
    ocr = build(tmp_path)
    m = _mod("detect_symbols")
    res, stats = m.process(tmp_path)
    page = res["TECH-DIAG_TEST_Exit_Lighting"]["pages"][0]["symbols"]
    assert len(page["emergency_light"]) == 5, page["emergency_light"]      # incl. 2 turned; legend not counted
    assert len(page["exit_sign"]) == 3
    assert all(b["box"][0] < 1900 for b in page["emergency_light"] + page["exit_sign"])   # nothing in the legend
    assert stats["exit_emergency_lighting"]["templates_found"] == ["emergency_light", "exit_sign"]
    assert (ocr / "symbols_review.html").exists()


def test_draft_counts_emergency_lights_from_symbols_with_that_basis(tmp_path, monkeypatch):
    build(tmp_path)
    d = _mod("draft_power_links")
    monkeypatch.setattr(d, "load_assets", lambda: {})
    assert d.main(["x", str(tmp_path)]) == 0
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "power_links_draft.xlsx")["Links"]
    head = [c.value for c in ws[1]]
    rows = [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    em = [r for r in rows if r["feeds"] == "emergency_light"]
    assert sum(r["count"] for r in em) == 5 and all(r["floor"] == "B1" for r in em)
    assert all("legend symbols found on the drawing" in r["basis"] for r in em)


def test_a_bad_legend_cut_is_reported_not_used(tmp_path):
    m = _mod("detect_symbols")
    blank = np.full((400, 600), 255, np.uint8)
    tpl, why = m.cut_symbol(blank, [300, 100, 500, 125])
    assert tpl is None and "no ink" in why


def test_noisy_blurred_scan_with_smaller_symbols_still_counts_correctly(tmp_path):
    ocr = build(tmp_path)
    img = cv2.imread(str(ocr / "TEST_p1.png"), cv2.IMREAD_GRAYSCALE)
    plan = cv2.resize(img[:1400, :1900], None, fx=0.9, fy=0.9)            # plan drawn 10 % smaller than the legend
    img[:1400, :1900] = 255
    img[:plan.shape[0], :plan.shape[1]] = plan
    img = cv2.GaussianBlur(img, (5, 5), 0)
    rng = np.random.default_rng(7)
    img = np.clip(img.astype(int) + rng.normal(0, 40, img.shape), 0, 255).astype(np.uint8)
    cv2.imwrite(str(ocr / "TEST_p1.png"), img)
    res, _ = _mod("detect_symbols").process(tmp_path)
    page = res["TECH-DIAG_TEST_Exit_Lighting"]["pages"][0]["symbols"]
    assert (len(page["emergency_light"]), len(page["exit_sign"])) == (5, 3)
