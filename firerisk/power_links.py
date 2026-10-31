"""Which fire-safety equipment each electrical panel supplies, from the building's drawings. ORIGIN: NEW (2026-10-31).

Source: `data/private/drawings/power_links.xlsx`, sheet "Links". The draft is made on the author's computer from the
local OCR of the scanned drawings (`scripts/draft_power_links.py`); the AUTHOR checks every row against the drawings
and writes Y in "confirmed". ONLY confirmed rows are used - an unconfirmed or invalid row is never used, because a
wrong link would give a wrong risk assessment. Invalid confirmed rows are reported, not guessed.

What the rules do with confirmed links (deterministic, documented in RISK_RULES.md, rule EQ-D1):
  * an asset that supplies life-safety equipment (emergency lights, exit signs, fire pump, smoke control, fire alarm)
    directly or through downstream panels counts as critical equipment (skill §2C.A: fire-safety equipment is
    critical) - so an open or repeated fault on it is judged as on critical equipment;
  * its alert states the consequence (which floors, which equipment types, how many) and adds a check of the
    emergency lighting / exit signs on battery;
  * it is linked to the fire-risk alert of every floor it supplies, with that reason.

What leaves the computer (data owner's condition 2): only the simplified summary from `summary()` - floor, equipment
type, count, direct or through a downstream panel. Panel labels, circuit codes, pages and notes from the drawings
stay private and are never put into an AI request, an API response for demo mode or a log.
"""

from __future__ import annotations

import os
from dataclasses import dataclass, field
from pathlib import Path

ROOT_DIR = Path(__file__).resolve().parent.parent
LINKS_FILE = Path(os.environ.get("FIRERISK_POWER_LINKS", ROOT_DIR / "data" / "private" / "drawings" / "power_links.xlsx"))
SHEET = "Links"
COLUMNS = ["confirmed", "floor", "panel_label", "asset_id", "feeds", "count", "fed_panel_label", "circuits",
           "source", "basis", "ocr_confidence", "note"]
FLOORS = ["B1", "F1", "F2", "F3", "F4", "Roof", "Building"]
FEEDS = {"emergency_light": "emergency lights", "exit_sign": "exit signs", "fire_pump": "fire pump",
         "smoke_control": "smoke control fans", "fire_alarm": "fire alarm panel", "other_fire_safety":
         "other fire-safety equipment", "sub_panel": "downstream panel"}
LIFE_SAFETY = {k for k in FEEDS if k != "sub_panel"}
YES = {"y", "yes", "x", "1", "true", "co", "có", "ok"}


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def norm_label(v) -> str:
    """Panel labels are compared without spaces, case or separators ('DB - B1' == 'db_b1')."""
    return "".join(ch for ch in _s(v).upper() if ch.isalnum())


@dataclass
class Link:
    row: int
    floor: str
    panel: str            # normalised panel label (private)
    asset_id: str | None
    feeds: str
    count: int | None
    fed_panel: str        # normalised, only for feeds == sub_panel


@dataclass
class PowerLinks:
    links: list[Link] = field(default_factory=list)
    problems: list[dict] = field(default_factory=list)      # {row, problem} - rows NOT used
    rows_total: int = 0
    rows_unconfirmed: int = 0
    present: bool = False
    error: str | None = None

    # --- lookups -------------------------------------------------------------------------------
    def panels_of_asset(self, asset_id: str) -> set[str]:
        return {l.panel for l in self.links if l.asset_id and l.asset_id == asset_id and l.panel}

    def _downstream(self, panels: set[str]) -> tuple[set[str], set[str]]:
        """(all panels reached through sub_panel links, panels reached only through another panel). Cycle-safe."""
        seen, todo = set(panels), list(panels)
        while todo:
            p = todo.pop()
            for l in self.links:
                if l.feeds == "sub_panel" and l.panel == p and l.fed_panel and l.fed_panel not in seen:
                    seen.add(l.fed_panel)
                    todo.append(l.fed_panel)
        return seen, seen - panels

    def dependents(self, asset_id: str | None) -> list[dict]:
        """Life-safety equipment supplied by this asset, directly or through downstream panels (simplified: floor,
        type, count, via). Empty when the asset has no confirmed link."""
        if not asset_id:
            return []
        own_panels = self.panels_of_asset(asset_id)
        all_panels, indirect = self._downstream(own_panels) if own_panels else (set(), set())
        out: dict[tuple, dict] = {}
        for l in self.links:
            if l.feeds not in LIFE_SAFETY:
                continue
            direct = (l.asset_id == asset_id) or (l.panel and l.panel in own_panels)
            via_down = bool(l.panel) and l.panel in indirect
            if not (direct or via_down):
                continue
            key = (l.floor, l.feeds, "direct" if direct else "through a downstream panel")
            d = out.setdefault(key, {"floor": l.floor, "equipment": l.feeds, "equipment_label": FEEDS[l.feeds],
                                     "count": 0, "count_known": True, "via": key[2]})
            if l.count is None:
                d["count_known"] = False
            else:
                d["count"] += l.count
        return sorted(out.values(), key=lambda d: (FLOORS.index(d["floor"]), d["equipment"], d["via"]))

    def assets_supplying(self, floor: str | None = None) -> list[str]:
        ids = {l.asset_id for l in self.links if l.asset_id}
        return sorted(a for a in ids if any(not floor or d["floor"] == floor for d in self.dependents(a)))

    def status(self) -> dict:
        return {"present": self.present, "error": self.error, "rows": self.rows_total,
                "confirmed_used": len(self.links), "unconfirmed": self.rows_unconfirmed,
                "rejected": len(self.problems), "assets_with_links": len({l.asset_id for l in self.links if l.asset_id}),
                "file": "data/private/drawings/power_links.xlsx"}


def describe(deps: list[dict]) -> str:
    """'B1: 6 emergency lights, 2 exit signs; F1: exit signs (number not recorded)' - simplified, no labels."""
    by_floor: dict[str, list[str]] = {}
    for d in deps:
        n = f"{d['count']} {d['equipment_label']}" if d["count_known"] else f"{d['equipment_label']} (number not recorded)"
        if d["via"] != "direct":
            n += " via a downstream panel"
        by_floor.setdefault(d["floor"], []).append(n)
    return "; ".join(f"{f}: {', '.join(v)}" for f, v in by_floor.items())


def summary(links: PowerLinks, asset_id: str | None = None, floor: str | None = None) -> dict:
    """The ONLY form of the drawing data that may reach an AI request (condition 2)."""
    assets = [asset_id] if asset_id else links.assets_supplying(floor)
    rows = []
    for a in assets:
        deps = [d for d in links.dependents(a) if not floor or d["floor"] == floor]
        if deps:
            rows.append({"asset_id": a, "supplies": [{k: d[k] for k in ("floor", "equipment", "count", "count_known", "via")}
                                                      for d in deps]})
    return {"links": rows, "basis": "building drawings, each link checked by a person against the drawing",
            "note": "Shows which life-safety equipment loses normal supply if this asset fails. Battery backup of the "
                    "equipment is not recorded."}


def load(path: Path | None = None, known_assets: set[str] | None = None) -> PowerLinks:
    """Read the confirmed links. Never raises: a missing or broken file gives an empty result with the reason."""
    path = Path(path or LINKS_FILE)
    res = PowerLinks()
    if not path.exists():
        return res
    res.present = True
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        if SHEET not in wb.sheetnames:
            res.error = f"sheet '{SHEET}' not found"
            return res
        rows = list(wb[SHEET].iter_rows(values_only=True))
        wb.close()
    except Exception as exc:                           # locked by Excel, corrupt, wrong format
        res.error = f"{type(exc).__name__}: file could not be read"
        return res
    if not rows:
        res.error = "empty sheet"
        return res
    head = [_s(h).lower() for h in rows[0]]
    col = {name: head.index(name) for name in COLUMNS if name in head}
    missing = [c for c in ("confirmed", "floor", "panel_label", "asset_id", "feeds", "count") if c not in col]
    if missing:
        res.error = f"columns missing: {', '.join(missing)}"
        return res
    get = lambda r, k: r[col[k]] if k in col and col[k] < len(r) else None
    for i, r in enumerate(rows[1:], start=2):
        if not any(_s(v) for v in r):
            continue
        res.rows_total += 1
        if _s(get(r, "confirmed")).lower() not in YES:
            res.rows_unconfirmed += 1
            continue
        bad = []
        floor = next((f for f in FLOORS if f.lower() == _s(get(r, "floor")).lower()), None)
        if not floor:
            bad.append(f"floor must be one of {', '.join(FLOORS)}")
        feeds = _s(get(r, "feeds")).lower()
        if feeds not in FEEDS:
            bad.append(f"feeds must be one of {', '.join(FEEDS)}")
        raw_count = get(r, "count")
        count = None
        if _s(raw_count):
            try:
                count = int(float(_s(raw_count)))
                if count < 0 or count != float(_s(raw_count)):
                    raise ValueError
            except ValueError:
                bad.append("count must be a whole number >= 0 (or empty when not known)")
        panel = norm_label(get(r, "panel_label"))
        asset = _s(get(r, "asset_id")) or None
        if not panel and not asset:
            bad.append("give the panel label or the asset_id")
        if asset and known_assets is not None and asset not in known_assets:
            bad.append(f"asset_id {asset} is not in the asset register")
        fed = norm_label(get(r, "fed_panel_label"))
        if feeds == "sub_panel" and not fed:
            bad.append("feeds = sub_panel needs fed_panel_label")
        if feeds == "sub_panel" and not panel:
            bad.append("feeds = sub_panel needs the supplying panel_label")
        if feeds == "sub_panel" and fed and fed == panel:
            bad.append("a panel cannot feed itself")
        if bad:
            res.problems.append({"row": i, "problem": "; ".join(bad)})
            continue
        res.links.append(Link(row=i, floor=floor, panel=panel, asset_id=asset, feeds=feeds, count=count,
                              fed_panel=fed if feeds == "sub_panel" else ""))
    # one panel label mapped to two different assets is a contradiction: drop those rows rather than guess
    owner: dict[str, set] = {}
    for l in res.links:
        if l.panel and l.asset_id:
            owner.setdefault(l.panel, set()).add(l.asset_id)
    clash = {p for p, a in owner.items() if len(a) > 1}
    if clash:
        for l in [l for l in res.links if l.panel in clash]:
            res.problems.append({"row": l.row, "problem": "this panel label is linked to two different asset_ids"})
        res.links = [l for l in res.links if l.panel not in clash]
    # carry the asset of a panel to its other rows (a panel's asset_id needs to be written once only)
    for l in res.links:
        if l.panel and not l.asset_id and l.panel in owner and len(owner[l.panel]) == 1:
            l.asset_id = next(iter(owner[l.panel]))
    return res


_CACHE: dict = {}


def current(known_assets: set[str] | None = None) -> PowerLinks:
    """Loaded once per file modification time."""
    try:
        mtime = LINKS_FILE.stat().st_mtime
    except OSError:
        return PowerLinks()
    key = (str(LINKS_FILE), mtime, frozenset(known_assets) if known_assets is not None else None)
    if _CACHE.get("key") != key:
        _CACHE["key"], _CACHE["value"] = key, load(LINKS_FILE, known_assets)
    return _CACHE["value"]
