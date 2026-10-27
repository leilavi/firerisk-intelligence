"""Read spreadsheet tabs into row dictionaries that keep their original row numbers.

ORIGIN: PORTED (skill §2.2 "read every tab, record row counts, detect truncation") + NEW code.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl


@dataclass
class TabData:
    sheet: str
    header: list[str]
    rows: list[tuple[int, dict]]   # (excel_row_number, {column: value})
    header_row: int

    @property
    def count(self) -> int:
        return len(self.rows)


def _is_blank(row) -> bool:
    return all(c is None or (isinstance(c, str) and not c.strip()) for c in row)


def read_tab(path: Path, sheet: str, header_row: int = 1) -> TabData:
    """Read one tab; header on `header_row` (1-based). Blank rows are skipped but numbering is kept."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[sheet]
        header: list[str] = []
        rows: list[tuple[int, dict]] = []
        for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if idx < header_row:
                continue
            if idx == header_row:
                header = [str(c).strip() if c is not None else f"col{i}" for i, c in enumerate(row)]
                continue
            if _is_blank(row):
                continue
            rows.append((idx, dict(zip(header, row))))
        return TabData(sheet=sheet, header=header, rows=rows, header_row=header_row)
    finally:
        wb.close()


def list_tabs(path: Path) -> list[str]:
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def suspicious_row_count(n: int) -> bool:
    """Skill §2B.9: a tab that stops at a round number may have been cut off by an export limit."""
    return n > 0 and n in {100, 200, 250, 500, 1000, 2000, 5000, 10000}
